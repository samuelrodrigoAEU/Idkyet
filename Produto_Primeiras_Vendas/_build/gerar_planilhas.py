"""Gera as duas planilhas do projeto:
- Produto_Principal/Ferramentas-Primeiras-Vendas.xlsx (produto principal)
- Order_Bump_03_Painel/Painel-das-Primeiras-Vendas.xlsx (order bump 3)

Uso: python3 gerar_planilhas.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
TINTA, PAPEL, SINAL, AZUL, VERDE, CINZA, LINHA = "161A22", "F5F2EB", "E8552A", "2D4FD6", "2E7D5B", "5E6370", "D8D1C2"
ENTRADA = PatternFill("solid", fgColor="FFF8EC")   # células para preencher
CALC = PatternFill("solid", fgColor="EEF0F6")      # células calculadas
CAB = PatternFill("solid", fgColor=TINTA)
fio = Side(style="thin", color=LINHA)
BORDA = Border(left=fio, right=fio, top=fio, bottom=fio)
QUEBRA = Alignment(wrap_text=True, vertical="top")


def titulo(ws, texto, sub=None):
    ws["A1"] = texto
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=TINTA)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name="Arial", size=10, color=CINZA)
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_view.showGridLines = False


def cabecalho(ws, linha, nomes, larguras=None):
    for i, nome in enumerate(nomes, start=1):
        c = ws.cell(row=linha, column=i, value=nome)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = CAB
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDA
    if larguras:
        for i, w in enumerate(larguras, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.row_dimensions[linha].height = 30


def celula(ws, ref, valor=None, tipo="entrada", fmt=None, negrito=False):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    c.fill = ENTRADA if tipo == "entrada" else CALC if tipo == "calc" else PatternFill()
    c.border = BORDA
    c.alignment = QUEBRA
    c.font = Font(name="Arial", size=10, bold=negrito)
    if fmt:
        c.number_format = fmt
    return c


def lista(ws, intervalo, opcoes):
    dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(intervalo)


def texto_corrido(ws, linha, paragrafos, col_larg=100):
    ws.column_dimensions["A"].width = col_larg
    for p in paragrafos:
        c = ws.cell(row=linha, column=1, value=p)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Arial", size=10, bold=p.isupper() or p.endswith(":"))
        ws.row_dimensions[linha].height = max(15, 15 * (len(p) // 95 + 1))
        linha += 1
    return linha


# ---------------------------------------------------------------- Ferramentas (produto principal)
def ferramentas():
    wb = Workbook()
    ws = wb.active
    ws.title = "Comece aqui"
    titulo(ws, "Ferramentas · Primeiras Vendas na Prática")
    texto_corrido(ws, 3, [
        "COMO USAR ESTA PLANILHA",
        "Células em bege são para você preencher. Células em azul-claro fazem as contas sozinhas: não digite nelas.",
        "Siga as abas da esquerda para a direita, na mesma ordem dos módulos do guia:",
        "1. Ponto de partida (Módulo 1) · 2. Matriz de ideias (Módulo 2) · 3. Ficha de oferta (Módulo 3) · 4. Registro de contatos (Módulos 4 e 6) · 5. Plano de 7 dias (Módulo 7) · 6. Revisão semanal (Módulo 7).",
        "Funciona no Excel e no Google Planilhas (Arquivo > Importar). Se usar no celular, prefira o app do Google Planilhas.",
        "Os exemplos do guia são simulações. Os números desta planilha são os seus: nenhum resultado vem preenchido.",
    ])

    # Ponto de partida
    ws = wb.create_sheet("Ponto de partida")
    titulo(ws, "Ponto de partida", "Módulo 1. Seja conservador: anote o que você tem de verdade hoje.")
    campos = [
        ("Data de início do plano", "Use o formato dd/mm/aaaa. A revisão semanal conta as semanas a partir daqui."),
        ("Minutos por dia disponíveis", "Número. Ex.: 45"),
        ("Habilidades que já uso", "O que fazem você ser chamado para ajudar"),
        ("Acessos (grupos, profissão, bairro, contatos)", "Onde estão pessoas que posso alcançar"),
        ("Dinheiro que aceito perder num teste (R$)", "Pode ser zero"),
        ("Restrições", "Horários, não aparecer em vídeo, não receber em casa..."),
        ("O tipo de venda mais fácil de testar primeiro parece ser...", "Sua intuição agora; o Módulo 2 testa"),
    ]
    cabecalho(ws, 4, ["Campo", "Sua resposta", "Dica"], [42, 50, 48])
    for i, (campo, dica) in enumerate(campos, start=5):
        celula(ws, f"A{i}", campo, tipo=None, negrito=True)
        celula(ws, f"B{i}")
        celula(ws, f"C{i}", dica, tipo=None)
        ws.row_dimensions[i].height = 42
    ws["B5"].number_format = "dd/mm/yyyy"

    # Matriz de ideias
    ws = wb.create_sheet("Matriz de ideias")
    titulo(ws, "Matriz de ideias", "Módulo 2. Notas de 1 (ruim) a 5 (ótimo). Toda nota 4 ou 5 precisa de uma evidência; sem evidência, a nota máxima é 3.")
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 30
    cabecalho(ws, 4, ["Critério", "Peso", "Ideia A: nota", "Evidência A", "Ideia B: nota", "Evidência B", "Ideia C: nota", "Evidência C"],
              [26, 7, 12, 30, 12, 30, 12, 30])
    celula(ws, "A5", "Nome da ideia", tipo=None, negrito=True)
    for col in "CEG":
        celula(ws, f"{col}5")
    ws.merge_cells("C5:D5"); ws.merge_cells("E5:F5"); ws.merge_cells("G5:H5")
    criterios = [("Demanda", 3), ("Margem", 2), ("Concorrência", 2), ("Dificuldade operacional", 2), ("Investimento inicial", 1)]
    for i, (nome, peso) in enumerate(criterios, start=6):
        celula(ws, f"A{i}", nome, tipo=None, negrito=True)
        celula(ws, f"B{i}", peso, tipo="calc")
        for col in "CDEFGH":
            celula(ws, f"{col}{i}")
        ws.row_dimensions[i].height = 32
    lista(ws, "C6:C10", ["1", "2", "3", "4", "5"])
    lista(ws, "E6:E10", ["1", "2", "3", "4", "5"])
    lista(ws, "G6:G10", ["1", "2", "3", "4", "5"])
    celula(ws, "A11", "Total (máx. 50)", tipo=None, negrito=True)
    celula(ws, "A12", "Decisão", tipo=None, negrito=True)
    for col, ev in (("C", "D"), ("E", "F"), ("G", "H")):
        celula(ws, f"{col}11", f"=SUMPRODUCT($B$6:$B$10,{col}6:{col}10)", tipo="calc", negrito=True)
        ws.merge_cells(f"{col}12:{ev}12")
        celula(ws, f"{col}12",
               f'=IF(COUNT({col}6:{col}10)<5,"Preencha as 5 notas",IF(AND({col}11>=35,COUNTIF({col}6:{col}10,1)=0),"Testar","Reformular ou descartar"))',
               tipo="calc", negrito=True)
    celula(ws, "A14", "Regra: testar se o total for 35 ou mais e nenhum critério tiver nota 1. Se duas ideias passarem, comece pela de maior demanda.", tipo=None)
    ws.merge_cells("A14:H14")
    ws.row_dimensions[14].height = 30

    # Ficha de oferta
    ws = wb.create_sheet("Ficha de oferta")
    titulo(ws, "Ficha de oferta", "Módulo 3. Responda às nove perguntas e calcule o preço mínimo.")
    cabecalho(ws, 4, ["Pergunta", "Sua resposta"], [34, 80])
    perguntas = ["O que é", "Para quem (por situação)", "Resolve o quê", "O que a pessoa recebe", "Como funciona",
                 "Prazo", "Preço e pagamento", "Garantia ou condições", "Próximo passo"]
    for i, p in enumerate(perguntas, start=5):
        celula(ws, f"A{i}", p, tipo=None, negrito=True)
        celula(ws, f"B{i}")
        ws.row_dimensions[i].height = 36
    cabecalho(ws, 15, ["Cálculo do preço mínimo", "Valor"])
    itens = [
        ("Custo direto por venda (R$)", None, "entrada", '"R$" #,##0.00'),
        ("Horas por entrega", None, "entrada", "0.0"),
        ("Valor da sua hora (R$)", None, "entrada", '"R$" #,##0.00'),
        ("Custo do tempo (R$)", "=B17*B18", "calc", '"R$" #,##0.00'),
        ("Taxa fixa do meio de pagamento (R$)", None, "entrada", '"R$" #,##0.00'),
        ("Taxa percentual do meio de pagamento (%)", None, "entrada", "0.00%"),
        ("Preço mínimo (R$)", "=IFERROR((B16+B19+B20)/(1-B21),\"\")", "calc", '"R$" #,##0.00'),
        ("Preço que você vai cobrar (R$)", None, "entrada", '"R$" #,##0.00'),
        ("Sobra por venda depois de custos e taxas (R$)", "=IF(B23=\"\",\"\",B23-B16-B20-B23*B21)", "calc", '"R$" #,##0.00'),
        ("Situação", "=IF(B23=\"\",\"Informe o preço\",IF(B23<B22,\"Abaixo do mínimo: revise\",\"Acima do mínimo\"))", "calc", None),
    ]
    for i, (nome, formula, tipo, fmt) in enumerate(itens, start=16):
        celula(ws, f"A{i}", nome, tipo=None, negrito=tipo == "calc")
        celula(ws, f"B{i}", formula, tipo=tipo, fmt=fmt)
    celula(ws, "A27", "A sobra inclui o pagamento do seu tempo dentro do preço. Taxas de plataformas mudam: confira a tabela atual no site oficial.", tipo=None)
    ws.merge_cells("A27:B27")

    # Registro de contatos
    ws = wb.create_sheet("Registro de contatos")
    titulo(ws, "Registro de contatos", "Módulos 4, 6 e 7. Uma linha por pessoa. Marque Sim conforme ela avança. A coluna Semana é calculada a partir da data de início.")
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 30
    cols = ["Data do contato", "Semana", "Nome", "Canal", "Sinal de interesse", "Respondeu?", "Conversa com diagnóstico?",
            "Proposta?", "Venda?", "Valor da venda (R$)", "Próxima ação", "Observação"]
    cabecalho(ws, 4, cols, [13, 9, 22, 16, 30, 11, 13, 11, 9, 13, 26, 30])
    for r in range(5, 205):
        for c in range(1, 13):
            cel = ws.cell(row=r, column=c)
            cel.border = BORDA
            cel.fill = CALC if c == 2 else ENTRADA
            cel.font = Font(name="Arial", size=10)
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=2, value=f"=IF(OR(A{r}=\"\",'Ponto de partida'!$B$5=\"\"),\"\",INT((A{r}-'Ponto de partida'!$B$5)/7)+1)")
        ws.cell(row=r, column=10).number_format = '"R$" #,##0.00'
    lista(ws, "D5:D204", ["Rede próxima", "Indicação", "Comunidade", "Contato direto", "TikTok", "Instagram", "WhatsApp", "Outro"])
    lista(ws, "F5:I204", ["Sim", "Não"])
    ws.freeze_panes = "C5"

    # Plano de 7 dias
    ws = wb.create_sheet("Plano de 7 dias")
    titulo(ws, "Plano de 7 dias", "Módulo 7. Marque o status de cada dia. Se atrasar, continue de onde parou.")
    cabecalho(ws, 4, ["Dia", "Tarefas", "Pronto quando", "Status", "Anotações"], [6, 52, 40, 14, 36])
    dias = [
        ("Ponto de partida, lista de 20 problemas, 3 ideias", "Aba Ponto de partida completa e 3 ideias escritas"),
        ("Matriz com evidências, ideia escolhida, método de validação", "Uma ideia com 35+ pontos e método escolhido"),
        ("Oferta em uma página, preço mínimo, teste de explicação", "Alguém de fora explicou a oferta corretamente"),
        ("Lista de 30 potenciais clientes, mensagens adaptadas, 10 envios", "10 mensagens enviadas e registradas"),
        ("2 conteúdos publicados, mais 10 envios, responder quem respondeu", "2 conteúdos no ar e 20 mensagens no total"),
        ("Diagnósticos e propostas, mais 10 envios, 1º acompanhamento", "Toda conversa aberta teve diagnóstico ou proposta"),
        ("Revisão semanal e um ajuste escolhido", "Revisão preenchida e um único ajuste definido"),
    ]
    for i, (tarefa, pronto) in enumerate(dias, start=5):
        celula(ws, f"A{i}", i - 4, tipo=None, negrito=True)
        celula(ws, f"B{i}", tarefa, tipo=None)
        celula(ws, f"C{i}", pronto, tipo=None)
        celula(ws, f"D{i}", "A fazer")
        celula(ws, f"E{i}")
        ws.row_dimensions[i].height = 34
    lista(ws, "D5:D11", ["A fazer", "Fazendo", "Feito"])
    cabecalho(ws, 14, ["Dia", "Conteúdo planejado (Módulo 5)", "Função", "Status", "Gerou perguntas ou mensagens?"])
    funcoes = ["Educar", "Demonstrar", "Educar", "Quebrar objeção", "Demonstrar", "Provar", "Convidar"]
    for i, f in enumerate(funcoes, start=15):
        celula(ws, f"A{i}", i - 14, tipo=None, negrito=True)
        celula(ws, f"B{i}")
        celula(ws, f"C{i}", f)
        celula(ws, f"D{i}", "A fazer")
        celula(ws, f"E{i}")
    lista(ws, "D15:D21", ["A fazer", "Publicado"])
    lista(ws, "E15:E21", ["Sim", "Não"])

    # Revisão semanal
    ws = wb.create_sheet("Revisão semanal")
    titulo(ws, "Revisão semanal", "Módulo 7. Os números vêm do Registro de contatos. Com menos de 30 contatos na semana, trate as taxas como pista, não como prova.")
    ws.merge_cells("A2:K2")
    ws.row_dimensions[2].height = 30
    cabecalho(ws, 4, ["Semana", "Contatados", "Responderam", "Conversas", "Propostas", "Vendas", "Taxa de resposta",
                      "Taxa de conversa", "Taxa de proposta", "Taxa de fechamento", "Etapa mais fraca"],
              [9, 12, 13, 11, 11, 9, 12, 12, 12, 13, 22])
    R = "'Registro de contatos'!"
    for i, sem in enumerate(range(1, 9), start=5):
        celula(ws, f"A{i}", sem, tipo=None, negrito=True)
        celula(ws, f"B{i}", f"=COUNTIF({R}$B$5:$B$204,A{i})", tipo="calc")
        for col, src in (("C", "F"), ("D", "G"), ("E", "H"), ("F", "I")):
            celula(ws, f"{col}{i}", f'=COUNTIFS({R}$B$5:$B$204,A{i},{R}${src}$5:${src}$204,"Sim")', tipo="calc")
        for col, num, den in (("G", "C", "B"), ("H", "D", "C"), ("I", "E", "D"), ("J", "F", "E")):
            celula(ws, f"{col}{i}", f'=IF({den}{i}=0,"",{num}{i}/{den}{i})', tipo="calc", fmt="0%")
        celula(ws, f"K{i}",
               f'=IF(B{i}=0,"",IF(C{i}=0,"Resposta",IF(D{i}=0,"Conversa",IF(E{i}=0,"Proposta",IF(F{i}=0,"Fechamento",'
               f'INDEX({{"Resposta","Conversa","Proposta","Fechamento"}},MATCH(MIN(G{i}:J{i}),G{i}:J{i},0)))))))', tipo="calc")
    cabecalho(ws, 15, ["Semana", "3 frases literais de potenciais clientes", "", "", "Objeções repetidas", "", "", "Conteúdo que gerou mensagens", "", "", "Único ajuste da próxima semana"])
    for i, sem in enumerate(range(1, 9), start=16):
        celula(ws, f"A{i}", sem, tipo=None, negrito=True)
        for a, b in (("B", "D"), ("E", "G"), ("H", "J")):
            ws.merge_cells(f"{a}{i}:{b}{i}")
            celula(ws, f"{a}{i}")
        celula(ws, f"K{i}")
        ws.row_dimensions[i].height = 40
    destino = RAIZ / "Produto_Principal" / "Ferramentas-Primeiras-Vendas.xlsx"
    wb.save(destino)
    return destino


# ---------------------------------------------------------------- Painel (order bump 3)
def painel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Leia primeiro"
    titulo(ws, "Painel das Primeiras Vendas")
    texto_corrido(ws, 3, [
        "O QUE ESTE PAINEL FAZ",
        "Mostra, sem você fazer contas, quantas pessoas avançam em cada etapa da venda, por semana e por canal; quanto dinheiro realmente sobra depois de custos e taxas; e qual etapa está travando.",
        "ORDEM DE USO:",
        "1. Configurações: preencha a data de início, suas taxas e a meta de atividade semanal.",
        "2. Contatos: registre cada pessoa (uma linha) e atualize a etapa sempre que ela avançar.",
        "3. Vendas e custos: lance cada venda e cada gasto (material, frete, impulsionamento, ferramentas).",
        "4. Tempo: anote os minutos dedicados por dia. Serve para calcular quanto vale sua hora de verdade.",
        "5. Painel: leia os números e o diagnóstico automático toda semana.",
        "Células em bege: você preenche. Células em azul-claro: calculadas. Não apague as fórmulas.",
        "IMPORTANTE:",
        "Visualizações, curtidas e seguidores não entram aqui de propósito. Eles medem alcance, não vendas. O painel mede pessoas que conversaram com você e dinheiro que entrou.",
        "Com poucos contatos (menos de 30 na semana), as taxas variam muito por acaso. Use o diagnóstico como pista para escolher o que testar, não como prova.",
        "Nenhum número vem preenchido. O arquivo Painel-EXEMPLO-dados-ficticios.xlsx mostra o painel funcionando com dados inventados, só para você ver como fica.",
    ])

    # Configurações
    ws = wb.create_sheet("Configurações")
    titulo(ws, "Configurações", "Preencha uma vez e atualize quando algo mudar.")
    cabecalho(ws, 4, ["Item", "Valor", "Observação"], [44, 18, 60])
    conf = [
        ("Data de início", None, "dd/mm/yyyy", "As semanas são contadas a partir desta data"),
        ("Taxa percentual do meio de pagamento principal", None, "0.00%", "Ex.: 8,99% numa plataforma de produto digital; 0% em Pix para pessoa física. Confira no site oficial"),
        ("Taxa fixa por venda (R$)", None, '"R$" #,##0.00', "Valor fixo cobrado por transação, se houver"),
        ("Imposto estimado sobre a receita (%)", None, "0.00%", "Se você é MEI, o imposto é fixo mensal: lance em Custos e deixe 0% aqui"),
        ("Meta semanal de novos contatos", 30, "0", "Meta de atividade, não de faturamento"),
        ("Meta semanal de conteúdos publicados", 3, "0", ""),
        ("Valor de referência da sua hora (R$)", None, '"R$" #,##0.00', "Usado para comparar com o que você realmente ganhou por hora"),
    ]
    for i, (nome, val, fmt, obs) in enumerate(conf, start=5):
        celula(ws, f"A{i}", nome, tipo=None, negrito=True)
        celula(ws, f"B{i}", val, fmt=fmt)
        celula(ws, f"C{i}", obs, tipo=None)
        ws.row_dimensions[i].height = 30
    cabecalho(ws, 14, ["Canais (edite os nomes se quiser)", "", ""])
    canais = ["Rede próxima", "Indicação", "Comunidade", "Contato direto", "TikTok", "Instagram", "WhatsApp", "Outro"]
    for i, c in enumerate(canais, start=15):
        celula(ws, f"A{i}", c)

    # Contatos
    ws = wb.create_sheet("Contatos")
    titulo(ws, "Contatos", "Uma linha por pessoa. Etapa = o ponto mais avançado a que a pessoa chegou. Perdido = disse não ou não respondeu aos dois acompanhamentos.")
    ws.merge_cells("A2:K2")
    ws.row_dimensions[2].height = 30
    cabecalho(ws, 4, ["Data do 1º contato", "Semana", "Nome", "Canal", "Origem (post, grupo, indicação de quem)", "Etapa atual",
                      "Nº da etapa", "Perdido?", "Motivo da perda", "Próxima ação", "Data da próxima ação"],
              [13, 9, 22, 16, 30, 16, 10, 10, 26, 28, 14])
    etapas = ["1 Contatado", "2 Respondeu", "3 Conversa", "4 Proposta", "5 Venda"]
    for r in range(5, 505):
        for c in range(1, 12):
            cel = ws.cell(row=r, column=c)
            cel.border = BORDA
            cel.fill = CALC if c in (2, 7) else ENTRADA
            cel.font = Font(name="Arial", size=10)
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=11).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=2, value=f'=IF(OR(A{r}="",Configurações!$B$5=""),"",INT((A{r}-Configurações!$B$5)/7)+1)')
        ws.cell(row=r, column=7, value=f'=IF(F{r}="","",VALUE(LEFT(F{r},1)))')
    lista(ws, "D5:D504", canais)
    lista(ws, "F5:F504", etapas)
    lista(ws, "H5:H504", ["Sim", "Não"])
    lista(ws, "I5:I504", ["Sem resposta", "Sem interesse", "Preço", "Timing", "Comprou de outro", "Não era o público", "Outro"])
    ws.freeze_panes = "D5"
    ws.conditional_formatting.add("A5:K504", FormulaRule(formula=['AND($K5<>"",$K5<=TODAY(),$H5<>"Sim",$G5<5)'],
                                                          fill=PatternFill("solid", fgColor="FBE3D9")))

    # Vendas e custos
    ws = wb.create_sheet("Vendas e custos")
    titulo(ws, "Vendas e custos", "Lance cada venda à esquerda e cada gasto à direita. As taxas são calculadas com o que está em Configurações.")
    ws.merge_cells("A2:N2")
    cabecalho(ws, 4, ["Data", "Semana", "Cliente", "Canal", "O que vendeu", "Valor (R$)", "Custo direto (R$)", "Taxas (R$)",
                      "Imposto (R$)", "Reembolsado?", "Sobra (R$)", "", "", ""],
              [12, 9, 20, 16, 24, 12, 13, 11, 11, 12, 12, 3, 12, 9])
    for r in range(5, 205):
        for c in range(1, 12):
            cel = ws.cell(row=r, column=c)
            cel.border = BORDA
            cel.fill = CALC if c in (2, 8, 9, 11) else ENTRADA
            cel.font = Font(name="Arial", size=10)
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=2, value=f'=IF(OR(A{r}="",Configurações!$B$5=""),"",INT((A{r}-Configurações!$B$5)/7)+1)')
        ws.cell(row=r, column=8, value=f'=IF(F{r}="","",F{r}*Configurações!$B$6+Configurações!$B$7)')
        ws.cell(row=r, column=9, value=f'=IF(F{r}="","",F{r}*Configurações!$B$8)')
        ws.cell(row=r, column=11, value=f'=IF(F{r}="","",IF(J{r}="Sim",-H{r}-G{r},F{r}-G{r}-H{r}-I{r}))')
        for c in (6, 7, 8, 9, 11):
            ws.cell(row=r, column=c).number_format = '"R$" #,##0.00'
    lista(ws, "D5:D204", canais)
    lista(ws, "J5:J204", ["Não", "Sim"])
    # custos gerais
    for i, (n, w) in enumerate([("Data", 12), ("Semana", 9), ("Gasto", 26), ("Categoria", 18), ("Valor (R$)", 12)], start=13):
        col = ws.cell(row=4, column=i).column_letter
        ws.column_dimensions[col].width = w
        c = ws.cell(row=4, column=i, value=n)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=SINAL)
        c.border = BORDA
    for r in range(5, 105):
        for c in range(13, 18):
            cel = ws.cell(row=r, column=c)
            cel.border = BORDA
            cel.fill = CALC if c == 14 else ENTRADA
        ws.cell(row=r, column=13).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=14, value=f'=IF(OR(M{r}="",Configurações!$B$5=""),"",INT((M{r}-Configurações!$B$5)/7)+1)')
        ws.cell(row=r, column=17).number_format = '"R$" #,##0.00'
    lista(ws, "P5:P104", ["Material", "Frete", "Ferramenta/assinatura", "Impulsionamento", "Imposto MEI", "Transporte", "Outro"])
    ws["L4"].fill = PatternFill()

    # Tempo
    ws = wb.create_sheet("Tempo")
    titulo(ws, "Tempo dedicado", "Anote os minutos por dia. Inclui conversas, conteúdo, produção e entrega.")
    cabecalho(ws, 4, ["Data", "Semana", "Minutos", "No que trabalhou"], [12, 9, 10, 60])
    for r in range(5, 125):
        for c in range(1, 5):
            cel = ws.cell(row=r, column=c)
            cel.border = BORDA
            cel.fill = CALC if c == 2 else ENTRADA
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=2, value=f'=IF(OR(A{r}="",Configurações!$B$5=""),"",INT((A{r}-Configurações!$B$5)/7)+1)')

    # Painel
    ws = wb.create_sheet("Painel")
    titulo(ws, "Painel", "Atualiza sozinho. Leia de cima para baixo: atividade, funil, dinheiro, diagnóstico.")
    ws.merge_cells("A2:N2")
    C, V, T = "Contatos!", "'Vendas e custos'!", "Tempo!"
    cab = ["Semana", "Novos contatos", "Meta", "Responderam", "Conversas", "Propostas", "Vendas",
           "Resposta", "Conversa", "Proposta", "Fechamento", "Etapa mais fraca", "Horas", "Diagnóstico"]
    cabecalho(ws, 4, cab, [9, 11, 8, 12, 11, 11, 9, 10, 10, 10, 11, 16, 8, 60])
    for i, sem in enumerate(range(1, 13), start=5):
        celula(ws, f"A{i}", sem, tipo=None, negrito=True)
        celula(ws, f"B{i}", f"=COUNTIF({C}$B$5:$B$504,A{i})", tipo="calc")
        celula(ws, f"C{i}", "=Configurações!$B$9", tipo="calc")
        for col, n in (("D", 2), ("E", 3), ("F", 4), ("G", 5)):
            celula(ws, f"{col}{i}", f'=COUNTIFS({C}$B$5:$B$504,A{i},{C}$G$5:$G$504,">="&{n})', tipo="calc")
        for col, num, den in (("H", "D", "B"), ("I", "E", "D"), ("J", "F", "E"), ("K", "G", "F")):
            celula(ws, f"{col}{i}", f'=IF({den}{i}=0,"",{num}{i}/{den}{i})', tipo="calc", fmt="0%")
        celula(ws, f"L{i}",
               f'=IF(B{i}=0,"",IF(D{i}=0,"Resposta",IF(E{i}=0,"Conversa",IF(F{i}=0,"Proposta",IF(G{i}=0,"Fechamento",'
               f'INDEX({{"Resposta","Conversa","Proposta","Fechamento"}},MATCH(MIN(H{i}:K{i}),H{i}:K{i},0)))))))', tipo="calc")
        celula(ws, f"M{i}", f"=SUMIF({T}$B$5:$B$124,A{i},{T}$C$5:$C$124)/60", tipo="calc", fmt="0.0")
        celula(ws, f"N{i}",
               f'=IF(B{i}=0,"",IF(B{i}<C{i},"Atividade abaixo da meta: antes de mudar a oferta, aumente os contatos. ","")&'
               f'IF(L{i}="Resposta","Poucas respostas: revise público, canal e personalização da 1ª mensagem.",'
               f'IF(L{i}="Conversa","Respondem sem interesse: fale com quem já mostrou o problema e deixe o benefício mais claro.",'
               f'IF(L{i}="Proposta","Conversas não viram proposta: conecte a solução ao que a pessoa disse e pergunte o prazo.",'
               f'IF(L{i}="Fechamento","Propostas sem venda: mostre demonstração, ofereça garantia simples e facilite o pagamento.","")))))',
               tipo="calc")
        ws.row_dimensions[i].height = 42
    # totais de funil
    t = 18
    cabecalho(ws, t, ["Funil acumulado", "Pessoas", "% do topo", "", "", "", "", "", "", "", "", "", "", ""])
    for j, (nome, n) in enumerate([("Contatados", 1), ("Responderam", 2), ("Conversa", 3), ("Proposta", 4), ("Venda", 5)], start=t + 1):
        celula(ws, f"A{j}", nome, tipo=None, negrito=True)
        celula(ws, f"B{j}", f'=COUNTIF({C}$G$5:$G$504,">="&{n})', tipo="calc")
        celula(ws, f"C{j}", f'=IF($B${t+1}=0,"",B{j}/$B${t+1})', tipo="calc", fmt="0%")
    # por canal
    k = 26
    cabecalho(ws, k, ["Canal", "Contatos", "Responderam", "Vendas", "Taxa de resposta", "Receita (R$)", "", "", "", "", "", "", "", ""])
    for j in range(8):
        r = k + 1 + j
        celula(ws, f"A{r}", f"=Configurações!A{15 + j}", tipo="calc", negrito=True)
        celula(ws, f"B{r}", f"=COUNTIF({C}$D$5:$D$504,A{r})", tipo="calc")
        celula(ws, f"C{r}", f'=COUNTIFS({C}$D$5:$D$504,A{r},{C}$G$5:$G$504,">=2")', tipo="calc")
        celula(ws, f"D{r}", f'=COUNTIFS({C}$D$5:$D$504,A{r},{C}$G$5:$G$504,5)', tipo="calc")
        celula(ws, f"E{r}", f'=IF(B{r}=0,"",C{r}/B{r})', tipo="calc", fmt="0%")
        celula(ws, f"F{r}", f'=SUMIFS({V}$F$5:$F$204,{V}$D$5:$D$204,A{r},{V}$J$5:$J$204,"<>Sim")', tipo="calc", fmt='"R$" #,##0.00')
    # dinheiro
    d = 37
    cabecalho(ws, d, ["Dinheiro (acumulado)", "Valor", "", "", "", "", "", "", "", "", "", "", "", ""])
    fin = [
        ("Receita bruta (sem reembolsos)", f'=SUMIFS({V}$F$5:$F$204,{V}$J$5:$J$204,"<>Sim")'),
        ("Custos diretos das vendas", f"=SUM({V}$G$5:$G$204)"),
        ("Taxas pagas", f"=SUM({V}$H$5:$H$204)"),
        ("Impostos estimados", f'=SUMIFS({V}$I$5:$I$204,{V}$J$5:$J$204,"<>Sim")'),
        ("Gastos gerais", f"=SUM({V}$Q$5:$Q$104)"),
        ("Resultado (o que sobrou)", f"=B{d+1}-B{d+2}-B{d+3}-B{d+4}-B{d+5}"),
        ("Vendas confirmadas", f'=COUNTIFS({V}$F$5:$F$204,">0",{V}$J$5:$J$204,"<>Sim")'),
        ("Reembolsos", f'=COUNTIF({V}$J$5:$J$204,"Sim")'),
        ("Ticket médio", f'=IF(B{d+7}=0,"",B{d+1}/B{d+7})'),
        ("Horas dedicadas", f"=SUM({T}$C$5:$C$124)/60"),
        ("Quanto você ganhou por hora", f'=IF(B{d+10}=0,"",B{d+6}/B{d+10})'),
        ("Comparação com sua hora de referência", f'=IF(OR(B{d+11}="",Configurações!$B$11=""),"",IF(B{d+11}>=Configurações!$B$11,"Acima da referência","Abaixo da referência"))'),
    ]
    for j, (nome, f) in enumerate(fin, start=d + 1):
        celula(ws, f"A{j}", nome, tipo=None, negrito=True)
        fmt = '"R$" #,##0.00' if j not in (d + 7, d + 8, d + 10, d + 12) else ("0.0" if j == d + 10 else None)
        celula(ws, f"B{j}", f, tipo="calc", fmt=fmt)
    ch = BarChart()
    ch.type = "bar"
    ch.title = "Funil acumulado"
    ch.y_axis.title = None
    ch.add_data(Reference(ws, min_col=2, min_row=t, max_row=t + 5), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=t + 1, max_row=t + 5))
    ch.legend = None
    ch.height, ch.width = 7, 12
    ch.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, "E18")
    ws.freeze_panes = "B5"

    destino = RAIZ / "Order_Bump_03_Painel" / "Painel-das-Primeiras-Vendas.xlsx"
    wb.save(destino)
    return destino


def exemplo(origem):
    """Cópia do painel com dados fictícios, para o comprador ver como fica preenchido."""
    from datetime import date, timedelta
    from openpyxl import load_workbook
    import random
    random.seed(7)
    wb = load_workbook(origem)
    inicio = date(2026, 11, 2)
    cf = wb["Configurações"]
    cf["B5"], cf["B6"], cf["B7"], cf["B8"], cf["B11"] = inicio, 0, 0, 0, 15
    ct = wb["Contatos"]
    canais = ["Rede próxima", "Indicação", "Comunidade", "Contato direto", "Instagram", "TikTok"]
    etapas = ["1 Contatado", "2 Respondeu", "3 Conversa", "4 Proposta", "5 Venda"]
    pesos = {"Rede próxima": [3, 3, 2, 1, 1], "Indicação": [1, 2, 2, 2, 2], "Comunidade": [4, 3, 2, 1, 0],
             "Contato direto": [8, 3, 1, 1, 0], "Instagram": [5, 2, 1, 1, 0], "TikTok": [3, 3, 2, 1, 0]}
    vendas = []
    for i in range(70):
        r = 5 + i
        d = inicio + timedelta(days=i // 5)
        canal = random.choice(canais)
        etapa = random.choices(etapas, weights=pesos[canal])[0]
        ct[f"A{r}"], ct[f"C{r}"], ct[f"D{r}"], ct[f"F{r}"] = d, f"Pessoa fictícia {i + 1}", canal, etapa
        if etapa.startswith("5"):
            vendas.append((d, f"Pessoa fictícia {i + 1}", canal))
    vc = wb["Vendas e custos"]
    for j, (d, nome, canal) in enumerate(vendas):
        r = 5 + j
        vc[f"A{r}"], vc[f"C{r}"], vc[f"D{r}"], vc[f"E{r}"], vc[f"F{r}"], vc[f"G{r}"], vc[f"J{r}"] = d, nome, canal, "Serviço (simulação)", 180, 32, "Não"
    vc["M5"], vc["O5"], vc["P5"], vc["Q5"] = inicio, "Fundo branco para fotos", "Material", 25
    tp = wb["Tempo"]
    for k in range(14):
        tp[f"A{5 + k}"], tp[f"C{5 + k}"], tp[f"D{5 + k}"] = inicio + timedelta(days=k), random.choice([30, 45, 60]), "Contatos e conteúdo (simulação)"
    wb["Leia primeiro"]["A1"] = "EXEMPLO COM DADOS FICTÍCIOS · não é resultado real"
    destino = origem.with_name("Painel-EXEMPLO-dados-ficticios.xlsx")
    wb.save(destino)
    return destino


if __name__ == "__main__":
    p = painel()
    for arq in (ferramentas(), p, exemplo(p)):
        print("ok", arq.relative_to(RAIZ))
