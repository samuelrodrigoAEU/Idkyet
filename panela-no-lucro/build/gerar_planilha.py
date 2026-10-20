"""Gera produto/Planilha-Panela-no-Lucro.xlsx. Rode: python3 build/gerar_planilha.py"""
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

OUT = "produto/Planilha-Panela-no-Lucro.xlsx"
VERDE, MOSTARDA, CREME, TOMATE, CINZA = "1F4D3A", "E8A33D", "FBF6EC", "C0392B", "6B6B6B"
F = "Calibri"
h1 = Font(name=F, size=16, bold=True, color=VERDE)
hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
bold = Font(name=F, size=10, bold=True)
norm = Font(name=F, size=10)
note = Font(name=F, size=9, italic=True, color=CINZA)
fill_hdr = PatternFill("solid", fgColor=VERDE)
fill_in = PatternFill("solid", fgColor="FFF4D6")   # célula que VOCÊ preenche
fill_calc = PatternFill("solid", fgColor="EEF3F0")  # célula calculada (não mexa)
thin = Side(style="thin", color="D9D9D9")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
BRL = 'R$ #,##0.00'
PCT = '0.0%'
wb = Workbook()


def header(ws, row, cols, widths=None):
    for i, t in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.border = hdr, fill_hdr, box
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 42
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def cell(ws, ref, value, kind="calc", fmt=None):
    c = ws[ref]
    c.value = value
    c.font, c.border = norm, box
    c.fill = fill_in if kind == "in" else fill_calc
    if fmt:
        c.number_format = fmt
    return c


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = h1
    ws["A2"] = sub
    ws["A2"].font = note
    ws.sheet_properties.tabColor = VERDE


# ---------------- 0. COMECE AQUI ----------------
ws = wb.active
ws.title = "Comece Aqui"
title(ws, "Panela no Lucro · Planilha de Custo, Preço e Pedidos", "Funciona no Excel, no Google Planilhas e no LibreOffice. Os valores de exemplo são da cozinha fictícia 'Cozinha da Cláudia' e servem só para você ver tudo funcionando.")
passos = [
    ("COR DAS CÉLULAS", ""),
    ("Amarelo claro", "você preenche"),
    ("Verde claro", "a planilha calcula. Não apague as fórmulas."),
    ("", ""),
    ("ORDEM DE USO (veja o Guia, Dias 1 a 7)", ""),
    ("1. Config", "Seus números: valor da sua hora, margem, gás, combustível, taxas e custos fixos do mês."),
    ("2. Ingredientes", "Tudo o que você compra, com o preço que pagou e o tamanho do pacote."),
    ("3. Fichas", "Cada receita, linha por linha: ingrediente e quantidade usada."),
    ("4. Produtos", "Rendimento, embalagem e tempo. Mostra o custo real, o preço mínimo, o preço sugerido e a situação do seu preço atual."),
    ("5. Entrega", "Seus bairros e distâncias. Mostra o custo real de cada entrega e a taxa sugerida."),
    ("6. Pedidos", "Anote cada pedido. A planilha calcula taxa, custo e lucro de cada um."),
    ("7. Semana", "Quantas unidades vai produzir. Gera a lista de compras na aba Ingredientes (coluna 'Comprar p/ semana')."),
    ("8. Painel", "Escolha o mês: faturamento, lucro real, dinheiro que ficou com você e quanto você ganhou por hora."),
    ("", ""),
    ("PARA COMEÇAR DO ZERO", "Apague só as células AMARELAS de exemplo (Ingredientes, Fichas, Produtos, Entrega e Pedidos) e digite as suas. Não apague linhas inteiras."),
    ("IMPORTANTE", "O nome do produto na aba Produtos tem de ser IGUAL ao nome da receita na aba Fichas (mesma grafia)."),
]
for i, (a, b) in enumerate(passos, 4):
    ws.cell(row=i, column=1, value=a).font = bold
    ws.cell(row=i, column=2, value=b).font = norm
ws["A5"].fill, ws["A6"].fill = fill_in, fill_calc
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 110

# ---------------- 1. CONFIG ----------------
c = wb.create_sheet("Config")
title(c, "1. Configurações", "Preencha uma vez e revise todo mês, ou quando o gás ou o combustível mudarem de preço.")
cfg = [
    # linha, rótulo, valor, formato, tipo, ajuda
    (4, "Quanto vale 1 hora do seu trabalho (R$)", 15, BRL, "in", "Use pelo menos o que você ganharia por hora em outro trabalho. Com R$ 1.518/mês e 220 h, a hora sai a ~R$ 6,90. Muita gente usa entre R$ 12 e R$ 20."),
    (5, "Margem de lucro desejada (% do preço)", 0.2, PCT, "in", "É o que sobra para o NEGÓCIO depois de pagar tudo, inclusive a sua hora. Entre 15% e 30% é comum. Não passe de 50%."),
    (6, "Preço do botijão de gás (R$)", 125, BRL, "in", "Preço que você pagou no último botijão."),
    (7, "Horas de 1 boca acesa que o botijão dura", 60, '0', "in", "Para medir: anote o dia da troca e quantas horas de boca usou (forno = 3 bocas). Sem medir, use 60."),
    (8, "Custo do gás por hora de boca (R$)", "=C6/C7", BRL, "calc", ""),
    (9, "Preço do combustível (R$/litro)", 6.2, BRL, "in", "Para entregas próprias."),
    (10, "Consumo do veículo (km por litro)", 35, '0.0', "in", "Moto ~30–40; carro ~9–12."),
    (11, "Desgaste e manutenção por km (R$)", 0.15, BRL, "in", "Pneu, óleo, revisão. Moto ~R$ 0,10–0,20; carro ~R$ 0,30–0,50."),
    (12, "Custo real por km rodado (R$)", "=C9/C10+C11", BRL, "calc", ""),
    (14, "TAXAS DE RECEBIMENTO (% que fica com a operadora)", None, None, None, "Confira no extrato da sua maquininha ou do app."),
    (15, "Pix", 0, PCT, "in", ""),
    (16, "Dinheiro", 0, PCT, "in", ""),
    (17, "Cartão de débito", 0.019, PCT, "in", ""),
    (18, "Cartão de crédito", 0.039, PCT, "in", ""),
    (19, "App de delivery", 0.27, PCT, "in", "Comissão + taxa de pagamento. Confira no contrato do seu plano."),
    (20, "Taxa média usada no preço de balcão", 0.01, PCT, "in", "Se quase tudo é Pix, use 0% a 1%. Se muito é crédito, use ~3%."),
    (22, "CUSTOS FIXOS DO MÊS (o que você paga vendendo pouco ou muito)", None, None, None, ""),
    (23, "Aumento estimado na conta de luz", 40, BRL, "in", "Compare a conta de agora com a de antes de começar a vender."),
    (24, "Aumento estimado na conta de água", 20, BRL, "in", ""),
    (25, "Parte do celular/internet usada no negócio", 30, BRL, "in", ""),
    (26, "Guia mensal do MEI (DAS), se tiver", 82, BRL, "in", "Confira o valor do ano no Portal do Empreendedor. Se não for MEI, ponha 0."),
    (27, "Divulgação (impulsionamento, panfleto)", 0, BRL, "in", ""),
    (28, "Reposição de utensílios (panela, forma), por mês", 30, BRL, "in", "Ex.: R$ 360 por ano / 12."),
    (29, "Outros fixos", 0, BRL, "in", ""),
    (30, "TOTAL DE CUSTOS FIXOS DO MÊS", "=SUM(C23:C29)", BRL, "calc", ""),
    (31, "Faturamento que você espera no mês (R$)", 3000, BRL, "in", "Seja realista: use a média dos últimos meses. Se está começando, chute baixo."),
    (32, "Peso dos fixos no preço (% do faturamento)", "=IF(C31>0,C30/C31,0)", PCT, "calc", "Se passar de 15%, seus fixos estão pesados para o tanto que você vende."),
]
for r, lab, val, fmt, kind, ajuda in cfg:
    c.cell(row=r, column=2, value=lab).font = bold if kind is None or kind == "calc" else norm
    if kind:
        cell(c, f"C{r}", val, kind, fmt)
    c.cell(row=r, column=4, value=ajuda).font = note
c.column_dimensions["A"].width = 2
c.column_dimensions["B"].width = 52
c.column_dimensions["C"].width = 14
c.column_dimensions["D"].width = 95
c["B34"] = "Alerta:"
c["B34"].font = bold
c["C34"] = '=IF(C5+C20+C32>=0.9,"ERRO: margem + taxa + fixos passam de 90%. Reduza a margem ou os fixos.","OK")'
c["C34"].font = Font(name=F, size=10, bold=True, color=TOMATE)

# ---------------- 2. INGREDIENTES ----------------
g = wb.create_sheet("Ingredientes")
title(g, "2. Ingredientes, embalagens e insumos", "Preço de EXEMPLO (out/2026). Troque pelo preço que VOCÊ pagou. Use g para peso, ml para líquidos e un para unidades.")
header(g, 3, ["Item", "Unidade (g, ml, un)", "Quantidade do pacote", "Preço pago (R$)", "Custo por unidade (R$)", "Comprar p/ semana (na unidade)", "Pacotes a comprar"],
       [30, 12, 14, 14, 16, 18, 14])
ingr = [
    ("Arroz", "g", 5000, 27.90), ("Feijão carioca", "g", 1000, 8.49), ("Peito de frango", "g", 1000, 19.90),
    ("Carne moída", "g", 1000, 38.90), ("Óleo de soja", "ml", 900, 7.99), ("Alho", "g", 200, 6.50),
    ("Cebola", "g", 1000, 5.99), ("Sal", "g", 1000, 3.29), ("Tomate", "g", 1000, 7.99),
    ("Cenoura", "g", 1000, 4.99), ("Farinha de trigo", "g", 1000, 5.49), ("Leite integral", "ml", 1000, 5.29),
    ("Margarina", "g", 500, 9.90), ("Caldo de galinha (tablete)", "un", 6, 3.49), ("Farinha de rosca", "g", 500, 6.90),
    ("Ovo", "un", 30, 21.90), ("Requeijão", "g", 400, 10.90), ("Leite condensado", "g", 395, 6.99),
    ("Creme de leite", "g", 200, 3.79), ("Chocolate em pó 50%", "g", 200, 9.90), ("Chocolate granulado", "g", 500, 14.90),
    ("Forminha nº 4 (cento)", "un", 100, 5.90), ("Pote marmita 750 ml c/ tampa", "un", 25, 24.90),
    ("Saco/embalagem coxinha (cento)", "un", 100, 12.90), ("Caixa brigadeiro 4 un", "un", 20, 17.90), ("Etiqueta/adesivo", "un", 100, 12.00),
]
FIRST, LAST = 4, 203  # 200 linhas
for i in range(FIRST, LAST + 1):
    data = ingr[i - FIRST] if i - FIRST < len(ingr) else (None, None, None, None)
    for col, v in zip("ABCD", data):
        cell(g, f"{col}{i}", v, "in", BRL if col == "D" else None)
    cell(g, f"E{i}", f'=IF(AND(ISNUMBER(C{i}),C{i}>0,ISNUMBER(D{i})),D{i}/C{i},"")', fmt='R$ #,##0.0000')
    cell(g, f"F{i}", f'=IF(A{i}="","",SUMIF(Fichas!$B$4:$B$603,A{i},Fichas!$F$4:$F$603))', fmt='#,##0')
    cell(g, f"G{i}", f'=IF(OR(A{i}="",F{i}="",N(C{i})=0),"",ROUNDUP(F{i}/C{i},0))', fmt='0')
g.freeze_panes = "A4"

# ---------------- 3. FICHAS ----------------
fi = wb.create_sheet("Fichas")
title(fi, "3. Fichas técnicas (receitas)", "Uma linha por ingrediente. Escreva o nome da receita igual em todas as linhas dela. Quantidade na MESMA unidade da aba Ingredientes. Inclua só o que vai na receita; a embalagem por unidade vai na aba Produtos.")
header(fi, 3, ["Receita (= nome do produto)", "Ingrediente (escolha da lista)", "Quantidade usada", "Unidade", "Custo na receita (R$)", "Quantidade p/ a semana"],
       [32, 32, 14, 10, 16, 18])
fichas = [
    # Marmita de frango: receita para 10 marmitas
    ("Marmita frango grelhado", "Arroz", 1000), ("Marmita frango grelhado", "Feijão carioca", 500),
    ("Marmita frango grelhado", "Peito de frango", 1500), ("Marmita frango grelhado", "Óleo de soja", 120),
    ("Marmita frango grelhado", "Alho", 30), ("Marmita frango grelhado", "Cebola", 200), ("Marmita frango grelhado", "Sal", 30),
    ("Marmita frango grelhado", "Tomate", 400), ("Marmita frango grelhado", "Cenoura", 300),
    # Coxinha: receita para 50 coxinhas
    ("Coxinha de frango (unid.)", "Peito de frango", 1000), ("Coxinha de frango (unid.)", "Farinha de trigo", 1000),
    ("Coxinha de frango (unid.)", "Leite integral", 1000), ("Coxinha de frango (unid.)", "Margarina", 50),
    ("Coxinha de frango (unid.)", "Caldo de galinha (tablete)", 2), ("Coxinha de frango (unid.)", "Cebola", 150),
    ("Coxinha de frango (unid.)", "Alho", 20), ("Coxinha de frango (unid.)", "Requeijão", 200),
    ("Coxinha de frango (unid.)", "Ovo", 3), ("Coxinha de frango (unid.)", "Farinha de rosca", 400),
    ("Coxinha de frango (unid.)", "Óleo de soja", 900),
    # Brigadeiro: receita para 40 unidades
    ("Brigadeiro gourmet (unid.)", "Leite condensado", 790), ("Brigadeiro gourmet (unid.)", "Creme de leite", 200),
    ("Brigadeiro gourmet (unid.)", "Chocolate em pó 50%", 80), ("Brigadeiro gourmet (unid.)", "Margarina", 30),
    ("Brigadeiro gourmet (unid.)", "Chocolate granulado", 250), ("Brigadeiro gourmet (unid.)", "Forminha nº 4 (cento)", 40),
]
for i in range(4, 604):  # 600 linhas
    data = fichas[i - 4] if i - 4 < len(fichas) else (None, None, None)
    cell(fi, f"A{i}", data[0], "in")
    cell(fi, f"B{i}", data[1], "in")
    cell(fi, f"C{i}", data[2], "in", '#,##0.##')
    cell(fi, f"D{i}", f'=IF(B{i}="","",IFERROR(VLOOKUP(B{i},Ingredientes!$A$4:$E$203,2,0),"?"))')
    cell(fi, f"E{i}", f'=IF(B{i}="","",IFERROR(VLOOKUP(B{i},Ingredientes!$A$4:$E$203,5,0)*C{i},"ITEM NÃO CADASTRADO"))', fmt=BRL)
    cell(fi, f"F{i}", f'=IF(A{i}="","",IFERROR(C{i}*VLOOKUP(A{i},Semana!$A$5:$E$54,5,0),0))', fmt='#,##0')
fi.freeze_panes = "A4"
dv_ing = DataValidation(type="list", formula1="=Ingredientes!$A$4:$A$203", allow_blank=True)
fi.add_data_validation(dv_ing)
dv_ing.add("B4:B603")

# ---------------- 4. PRODUTOS ----------------
p = wb.create_sheet("Produtos")
title(p, "4. Produtos e preços", "Custo real por unidade vendida e três preços: MÍNIMO (paga tudo, inclusive a sua hora, sem lucro), SUGERIDO (com a sua margem) e no APP.")
cols = ["Produto (= nome da receita)", "Rende quantas unidades", "Custo dos ingredientes da receita", "Embalagem por unidade (R$)",
        "Minutos de boca acesa na receita", "Minutos de trabalho na receita", "Ingredientes por unidade", "Gás por unidade",
        "CUSTO DE MATERIAL por unidade", "Sua hora por unidade", "CUSTO TOTAL por unidade", "PREÇO MÍNIMO",
        "PREÇO SUGERIDO", "Preço sugerido no APP", "Seu preço atual (R$)", "Lucro por unidade no preço atual",
        "Dinheiro p/ você por unidade (hora + lucro)", "SITUAÇÃO"]
header(p, 3, cols, [30, 11, 13, 12, 12, 12, 12, 10, 13, 11, 13, 12, 12, 12, 12, 13, 15, 26])
prods = [
    ("Marmita frango grelhado", 10, None, 1.20, 90, 120, 22),
    ("Coxinha de frango (unid.)", 50, None, 0.26, 150, 240, 4),
    ("Brigadeiro gourmet (unid.)", 40, None, 0.10, 30, 120, 2.5),
]
den_min = "(1-Config!$C$20-Config!$C$32)"
den_sug = "(1-Config!$C$5-Config!$C$20-Config!$C$32)"
den_app = "(1-Config!$C$5-Config!$C$19-Config!$C$32)"
for i in range(4, 54):  # 50 produtos
    d = prods[i - 4] if i - 4 < len(prods) else (None,) * 7
    cell(p, f"A{i}", d[0], "in")
    cell(p, f"B{i}", d[1], "in", '0')
    cell(p, f"C{i}", f'=IF(A{i}="","",SUMIF(Fichas!$A$4:$A$603,A{i},Fichas!$E$4:$E$603))', fmt=BRL)
    cell(p, f"D{i}", d[3], "in", BRL)
    cell(p, f"E{i}", d[4], "in", '0')
    cell(p, f"F{i}", d[5], "in", '0')
    ok = f'AND(A{i}<>"",N(B{i})>0)'
    cell(p, f"G{i}", f'=IF({ok},C{i}/B{i},"")', fmt=BRL)
    cell(p, f"H{i}", f'=IF({ok},N(E{i})/60*Config!$C$8/B{i},"")', fmt=BRL)
    cell(p, f"I{i}", f'=IF({ok},G{i}+N(D{i})+H{i},"")', fmt=BRL)
    cell(p, f"J{i}", f'=IF({ok},N(F{i})/60*Config!$C$4/B{i},"")', fmt=BRL)
    cell(p, f"K{i}", f'=IF({ok},I{i}+J{i},"")', fmt=BRL)
    cell(p, f"L{i}", f'=IF({ok},ROUNDUP(K{i}/{den_min}*10,0)/10,"")', fmt=BRL)
    cell(p, f"M{i}", f'=IF({ok},ROUNDUP(K{i}/{den_sug}*2,0)/2,"")', fmt=BRL)
    cell(p, f"N{i}", f'=IF({ok},IF({den_app}<=0,"MARGEM ALTA DEMAIS",ROUNDUP(K{i}/{den_app}*2,0)/2),"")', fmt=BRL)
    cell(p, f"O{i}", d[6], "in", BRL)
    cell(p, f"P{i}", f'=IF(AND({ok},N(O{i})>0),O{i}*{den_min}-K{i},"")', fmt=BRL)
    cell(p, f"Q{i}", f'=IF(P{i}="","",P{i}+J{i})', fmt=BRL)
    cell(p, f"R{i}", f'=IF(P{i}="","",IF(O{i}<L{i},"PREJUÍZO: abaixo do mínimo",IF(O{i}<M{i},"Atenção: abaixo do sugerido","OK: preço saudável")))')
p.freeze_panes = "B4"
p.conditional_formatting.add("R4:R53", FormulaRule(formula=['LEFT(R4,3)="PRE"'], fill=PatternFill("solid", fgColor="F8D7DA"), font=Font(color=TOMATE, bold=True)))
p.conditional_formatting.add("R4:R53", FormulaRule(formula=['LEFT(R4,3)="Ate"'], fill=PatternFill("solid", fgColor="FFF3CD")))
p.conditional_formatting.add("R4:R53", FormulaRule(formula=['LEFT(R4,2)="OK"'], fill=PatternFill("solid", fgColor="D4EDDA")))
p["A55"] = "Como ler: se o SEU PREÇO estiver abaixo do MÍNIMO, cada unidade vendida paga menos que o seu custo (contando a sua hora). Entre o mínimo e o sugerido, você se paga, mas o negócio não guarda reserva."
p["A55"].font = note

# ---------------- 5. ENTREGA ----------------
e = wb.create_sheet("Entrega")
title(e, "5. Entrega por bairro", "Distância de IDA (veja no Google Maps). Se paga motoboy, preencha o valor pago e a planilha ignora o cálculo de km.")
header(e, 3, ["Bairro / região", "Km de ida", "Minutos ida + volta", "Pago a motoboy (R$)", "CUSTO REAL da entrega", "TAXA SUGERIDA"], [28, 10, 12, 14, 14, 14])
bairros = [("Mesmo bairro", 1, 15, None), ("Centro", 3.5, 25, None), ("Jardim América", 6, 35, None), ("Vila Nova (motoboy)", 8, None, 12)]
for i in range(4, 44):
    d = bairros[i - 4] if i - 4 < len(bairros) else (None,) * 4
    for col, v, fmt in zip("ABCD", d, (None, '0.0', '0', BRL)):
        cell(e, f"{col}{i}", v, "in", fmt)
    cell(e, f"E{i}", f'=IF(A{i}="","",IF(N(D{i})>0,D{i},N(B{i})*2*Config!$C$12+N(C{i})/60*Config!$C$4))', fmt=BRL)
    cell(e, f"F{i}", f'=IF(E{i}="","",MAX(2,ROUNDUP(E{i}*2,0)/2))', fmt=BRL)
e["A45"] = "A taxa sugerida cobre combustível, desgaste e o SEU tempo no trânsito (com o valor da sua hora). Arredonda para cima de R$ 0,50 em R$ 0,50, com mínimo de R$ 2."
e["A45"].font = note

# ---------------- 6. PEDIDOS ----------------
pe = wb.create_sheet("Pedidos")
title(pe, "6. Pedidos", "Uma linha por pedido. Preencha as colunas amarelas. Se o 'Preço unitário' ficar vazio, a planilha usa o 'Seu preço atual' da aba Produtos.")
pc = ["Data", "Cliente", "Produto", "Qtd", "Preço unitário cobrado (vazio = tabela)", "Bairro (entrega)", "Taxa de entrega cobrada", "Pagamento",
      "Status", "Sinal recebido (R$)", "TOTAL do pedido", "Taxa da operadora", "Custo de material", "Sua hora (custo)", "Custo da entrega", "LUCRO do pedido", "Falta receber", "Mês", "Ano"]
header(pe, 3, pc, [11, 18, 28, 6, 14, 20, 11, 13, 13, 11, 12, 11, 12, 11, 11, 12, 11, 6, 7])
hoje = date(2026, 10, 1)
peds = [
    (date(2026, 10, 2), "Ana (vizinha)", "Marmita frango grelhado", 5, None, "Mesmo bairro", 0, "Pix", "Entregue", 0),
    (date(2026, 10, 3), "Escritório Silva", "Marmita frango grelhado", 8, 21, "Centro", 6, "Pix", "Entregue", 0),
    (date(2026, 10, 5), "Festa Júlia", "Coxinha de frango (unid.)", 100, 1.8, "Jardim América", 8, "Pix", "Entregue", 90),
    (date(2026, 10, 6), "Marcos", "Marmita frango grelhado", 2, None, "Mesmo bairro", 3, "Cartão de crédito", "Entregue", 0),
    (date(2026, 10, 8), "Festa Júlia", "Brigadeiro gourmet (unid.)", 80, None, "Jardim América", 0, "Pix", "Entregue", 100),
    (date(2026, 10, 10), "Bruna", "Marmita frango grelhado", 10, 20, "Vila Nova (motoboy)", 12, "Pix", "Entregue", 0),
    (date(2026, 10, 12), "Igreja", "Coxinha de frango (unid.)", 200, 1.6, "Centro", 0, "Dinheiro", "Entregue", 160),
    (date(2026, 10, 15), "iFood", "Marmita frango grelhado", 6, 26, "", 0, "App de delivery", "Entregue", 0),
    (date(2026, 10, 17), "Carla", "Brigadeiro gourmet (unid.)", 40, None, "Centro", 6, "Cartão de débito", "Entregue", 0),
    (date(2026, 10, 22), "Aniversário Léo", "Coxinha de frango (unid.)", 150, None, "Jardim América", 8, "Pix", "Sinal pago", 150),
    (date(2026, 10, 23), "Escritório Silva", "Marmita frango grelhado", 8, 21, "Centro", 6, "Pix", "Encomenda", 0),
]
for i in range(4, 504):  # 500 pedidos
    d = peds[i - 4] if i - 4 < len(peds) else (None,) * 10
    for col, v, fmt in zip("ABCDEFGHIJ", d, ('dd/mm/yyyy', None, None, '0', BRL, None, BRL, None, None, BRL)):
        cell(pe, f"{col}{i}", v, "in", fmt)
    act = f'AND(C{i}<>"",N(D{i})>0,I{i}<>"Cancelado")'
    preco = f'IF(N(E{i})>0,E{i},IFERROR(VLOOKUP(C{i},Produtos!$A$4:$O$53,15,0),0))'
    cell(pe, f"K{i}", f'=IF({act},D{i}*{preco}+N(G{i}),"")', fmt=BRL)
    cell(pe, f"L{i}", f'=IF(K{i}="","",K{i}*IFERROR(VLOOKUP(H{i},Config!$B$15:$C$19,2,0),0))', fmt=BRL)
    cell(pe, f"M{i}", f'=IF(K{i}="","",D{i}*IFERROR(VLOOKUP(C{i},Produtos!$A$4:$I$53,9,0),0))', fmt=BRL)
    cell(pe, f"N{i}", f'=IF(K{i}="","",D{i}*IFERROR(VLOOKUP(C{i},Produtos!$A$4:$J$53,10,0),0))', fmt=BRL)
    cell(pe, f"O{i}", f'=IF(OR(K{i}="",F{i}=""),IF(K{i}="","",0),IFERROR(VLOOKUP(F{i},Entrega!$A$4:$E$43,5,0),0))', fmt=BRL)
    cell(pe, f"P{i}", f'=IF(K{i}="","",K{i}-L{i}-M{i}-N{i}-O{i})', fmt=BRL)
    cell(pe, f"Q{i}", f'=IF(K{i}="","",IF(I{i}="Entregue",0,K{i}-N(J{i})))', fmt=BRL)
    cell(pe, f"R{i}", f'=IF(ISNUMBER(A{i}),MONTH(A{i}),"")')
    cell(pe, f"S{i}", f'=IF(ISNUMBER(A{i}),YEAR(A{i}),"")')
pe.freeze_panes = "D4"
for rng, src in (("C4:C503", "=Produtos!$A$4:$A$53"), ("F4:F503", "=Entrega!$A$4:$A$43"),
                 ("H4:H503", '"Pix,Dinheiro,Cartão de débito,Cartão de crédito,App de delivery"'),
                 ("I4:I503", '"Encomenda,Sinal pago,Entregue,Cancelado"')):
    dv = DataValidation(type="list", formula1=src, allow_blank=True)
    pe.add_data_validation(dv)
    dv.add(rng)
pe.conditional_formatting.add("P4:P503", FormulaRule(formula=['AND(ISNUMBER(P4),P4<0)'], fill=PatternFill("solid", fgColor="F8D7DA"), font=Font(color=TOMATE, bold=True)))

# ---------------- 7. SEMANA ----------------
s = wb.create_sheet("Semana")
title(s, "7. Planejamento da semana e lista de compras", "Coluna B: encomendas em aberto (status Encomenda ou Sinal pago). Na coluna C, digite quantas unidades vai produzir (encomendas + pronta-entrega). A lista de compras aparece na aba Ingredientes, colunas F e G.")
header(s, 4, ["Produto", "Encomendas em aberto (unid.)", "Vou produzir (unid.)", "Rendimento da receita", "Vezes a receita"], [30, 14, 14, 12, 12])
s.row_dimensions[3].height = 6
for i in range(5, 55):
    cell(s, f"A{i}", f'=IF(Produtos!A{i-1}="","",Produtos!A{i-1})')
    cell(s, f"B{i}", f'=IF(A{i}="","",SUMIFS(Pedidos!$D$4:$D$503,Pedidos!$C$4:$C$503,A{i},Pedidos!$I$4:$I$503,"Encomenda")+SUMIFS(Pedidos!$D$4:$D$503,Pedidos!$C$4:$C$503,A{i},Pedidos!$I$4:$I$503,"Sinal pago"))', fmt='0')
    cell(s, f"C{i}", {5: 30, 6: 150, 7: 40}.get(i), "in", '0')
    cell(s, f"D{i}", f'=IF(A{i}="","",Produtos!B{i-1})', fmt='0')
    cell(s, f"E{i}", f'=IF(OR(A{i}="",N(D{i})=0),0,N(C{i})/D{i})', fmt='0.00')

# ---------------- 8. PAINEL ----------------
d = wb.create_sheet("Painel", 1)
title(d, "8. Painel do mês", "Escolha o mês e o ano (células amarelas). Considera só pedidos não cancelados. Encomendas futuras do mês entram como previsão.")
cell(d, "C4", 10, "in", '0'); d["B4"] = "Mês (1 a 12)"
cell(d, "C5", 2026, "in", '0'); d["B5"] = "Ano"
cell(d, "C6", 52, "in", '0'); d["B6"] = "Horas que você trabalhou no negócio no mês (cozinhar, comprar, entregar, atender)"
crit = 'Pedidos!$R$4:$R$503,$C$4,Pedidos!$S$4:$S$503,$C$5'
rows = [
    (8, "Pedidos no mês", f'=COUNTIFS({crit},Pedidos!$K$4:$K$503,">0")', '0'),
    (9, "Faturamento (tudo que o cliente pagou)", f'=SUMIFS(Pedidos!$K$4:$K$503,{crit})', BRL),
    (10, "(-) Taxas de operadora e app", f'=SUMIFS(Pedidos!$L$4:$L$503,{crit})', BRL),
    (11, "(-) Material (ingrediente, embalagem, gás)", f'=SUMIFS(Pedidos!$M$4:$M$503,{crit})', BRL),
    (12, "(-) Entregas", f'=SUMIFS(Pedidos!$O$4:$O$503,{crit})', BRL),
    (13, "(-) Custos fixos do mês (aba Config)", '=Config!$C$30', BRL),
    (14, "→ DINHEIRO QUE SOBROU PARA VOCÊ", '=C9-C10-C11-C12-C13', BRL),
    (15, "Sua hora já paga dentro disso (valor da hora × horas planejadas nas receitas)", f'=SUMIFS(Pedidos!$N$4:$N$503,{crit})', BRL),
    (16, "→ LUCRO DO NEGÓCIO (o que sobra depois de pagar a sua hora)", '=C14-C15', BRL),
    (17, "Ticket médio por pedido", '=IF(C8>0,C9/C8,0)', BRL),
    (18, "VOCÊ GANHOU POR HORA DE TRABALHO REAL", '=IF(N(C6)>0,C14/C6,0)', BRL),
    (19, "Meta: valor da sua hora (Config)", '=Config!$C$4', BRL),
    (20, "Diagnóstico", '=IF(C9=0,"Sem pedidos neste mês",IF(C14<0,"PREJUÍZO: você pagou para trabalhar. Revise preços (aba Produtos) e fixos.",IF(C18<C19,"Sobrou dinheiro, mas sua hora real ficou ABAIXO da meta. Veja o tempo gasto e o lucro por produto.","Saudável: a hora real está igual ou acima da meta.")))', None),
    (21, "Falta receber (encomendas e sinais)", f'=SUMIFS(Pedidos!$Q$4:$Q$503,{crit})', BRL),
]
for r, lab, f, fmt in rows:
    d.cell(row=r, column=2, value=lab).font = bold if lab.startswith("→") or lab.isupper() or "VOCÊ" in lab else norm
    cell(d, f"C{r}", f, fmt=fmt)
d["C20"].alignment = Alignment(wrap_text=True)
d.row_dimensions[20].height = 45
header_row = 23
for i, t in enumerate(["Produto", "Unidades vendidas", "Faturamento", "Lucro do negócio", "Lucro por unidade"], 2):
    x = d.cell(row=header_row, column=i, value=t)
    x.font, x.fill, x.border = hdr, fill_hdr, box
d.cell(row=22, column=2, value="Lucro por produto no mês (descubra o que vale a pena vender)").font = bold
for k in range(15):
    r = header_row + 1 + k
    pr = k + 4
    cell(d, f"B{r}", f'=IF(Produtos!A{pr}="","",Produtos!A{pr})')
    cell(d, f"C{r}", f'=IF(B{r}="","",SUMIFS(Pedidos!$D$4:$D$503,Pedidos!$C$4:$C$503,B{r},{crit}))', fmt='0')
    cell(d, f"D{r}", f'=IF(B{r}="","",SUMIFS(Pedidos!$K$4:$K$503,Pedidos!$C$4:$C$503,B{r},{crit}))', fmt=BRL)
    cell(d, f"E{r}", f'=IF(B{r}="","",SUMIFS(Pedidos!$P$4:$P$503,Pedidos!$C$4:$C$503,B{r},{crit}))', fmt=BRL)
    cell(d, f"F{r}", f'=IF(OR(B{r}="",N(C{r})=0),"",E{r}/C{r})', fmt=BRL)
d.column_dimensions["A"].width = 2
d.column_dimensions["B"].width = 70
for col in "CDEF":
    d.column_dimensions[col].width = 18
d["B40"] = "Obs.: o 'lucro do pedido' (aba Pedidos) não desconta os custos fixos. Eles são descontados uma vez por mês, aqui no painel."
d["B40"].font = note

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
wb.save(OUT)
print("ok", OUT)
